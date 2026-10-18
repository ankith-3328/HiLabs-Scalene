"""Centene price transparency pipeline (default scope: California commercial group PPO).

Writes to output/:
  1_ppo_rate_file.xlsx          C6  deliverable 1, the rate file: one row per NPI + TIN + address
                                    + billing code
  c1_index_plans.csv            C1  every plan in every Centene index file (duplicates removed)
  c2_ppo_plans.csv              C2  the in-scope plans of the target network type
  c3_file_quality.md            C3  quality metrics and scores for every file used
  c4_ppo_rates.csv              C4  NPI, TIN, billing code and the rate chosen for them
  c5_ppo_rates_by_address.csv   C5  the C4 rates with every linked address, its confidence and the specialty

    uv run main.py                                    # full run
    uv run main.py --index-only                       # C1/C2 (and index file grades) only
    uv run main.py --state OR --network HMO           # another scope, same code

The run's metrics are printed to the console. deliverables/ holds the submitted run's files
and is never written to.
"""

import argparse
import csv
import json
import os
import sys
import tempfile
from collections import Counter
from concurrent.futures import ProcessPoolExecutor, ThreadPoolExecutor
from datetime import date
from pathlib import Path
from urllib.parse import unquote, urlsplit

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from network_mapping.address_linking import cms_practice_sites, irs_organizations, link
from network_mapping.downloader import FetchError, fetch
from network_mapping.file_quality import Grader, grade_index, npi_valid, quality_report
from network_mapping.npi_registry import lookup_all
from network_mapping.plan_index import (
    US_STATES,
    index_links,
    network_type,
    parse_index,
    state,
)
from network_mapping.rate_file import (
    collector,
    flatten,
    is_outlier,
    merge,
    rate_rows,
    service_medians,
)
from network_mapping.specialty_lookup import load_nucc, load_supplement, specialty

PAGE_URL = "https://www.centene.com/price-transparency-files.html"
SUPPLEMENT = Path(__file__).parent / "network_mapping" / "data" / "hilabs_rate_benchmarking_weights.xlsx"
PLAN_COLUMNS = ["brand", "plan_name", "plan_id", "market_type", "rate_file_url"]
PROVIDER_RATE_COLUMNS = ["NPI", "TIN", "Billing Code", "Address", "Specialty", "Rate"]


def write_csv(path, columns, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path, columns, rows, numeric=()):
    """One sheet with a bold, frozen, filterable header; columns in numeric are written as numbers,
    the rest as text (so NPIs and codes keep their leading zeros)."""
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(Path(path).stem[:31])
    ws.freeze_panes = "A2"
    for i, c in enumerate(columns):
        ws.column_dimensions[get_column_letter(i + 1)].width = 60 if c == "Address" else max(14, len(c) + 4)
    header = []
    for c in columns:
        cell = WriteOnlyCell(ws, c)
        cell.font = Font(bold=True)
        header.append(cell)
    ws.append(header)
    for r in rows:
        ws.append([float(r[c]) if c in numeric and r[c] else r[c] for c in columns])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    wb.save(path)


def unique(rows, key):
    """The first row for each distinct value of the key columns, in first-seen order."""
    first = {}
    for r in rows:
        first.setdefault(tuple(r[c] for c in key), r)
    return list(first.values())


def load_plans(errors):
    """Download Centene's page and every index file it links to (in parallel).

    Returns (index URLs, plan rows, {index URL: quality metrics}).
    """
    urls = index_links(fetch(PAGE_URL).decode("utf-8", errors="replace"), PAGE_URL)
    if not urls:
        errors.append(f"{PAGE_URL}: no .json index links found")

    def one(url):
        try:
            doc = json.loads(fetch(url))
        except (FetchError, ValueError) as err:
            errors.append(f"{url}: {err}")
            return [], None
        return parse_index(doc, url), grade_index(doc)

    with ThreadPoolExecutor(8) as pool:
        results = list(pool.map(one, urls))
    return urls, [r for rows, _ in results for r in rows], {u: g for u, (_, g) in zip(urls, results, strict=True) if g}


def process_rate_file(url):
    """Download one rate file to a temporary folder (deleted afterwards), stream it,
    and collect rates and quality counters on the way."""
    groups, tin_names, grader = {}, {}, Grader()
    collect = collector(groups, tin_names)
    with tempfile.TemporaryDirectory() as tmp:
        path = fetch(url, Path(tmp) / "rate_file.json")
        stats = flatten(path, os.devnull, on_row=lambda parts: (collect(parts), grader.add(parts)))
    grader.seen = grader.seen_exact = None  # only needed while streaming
    return groups, tin_names, grader, stats


def share(part, whole):
    return f"{len(part):,} of {len(whole):,} ({len(part) / max(len(whole), 1):.1%})"


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--state", default="CA", help="two-letter state (default CA)")
    ap.add_argument("--market", default="group", help="plan_market_type to keep (default group)")
    ap.add_argument("--network", default="PPO", help="network type to map (default PPO)")
    ap.add_argument("--supplement", default=SUPPLEMENT, type=Path, help="HiLabs specialty weights workbook")
    ap.add_argument("--out", default="output", help="output folder (default output)")
    ap.add_argument("--index-only", action="store_true", help="stop after C1/C2")
    args = ap.parse_args(argv)
    net = args.network.lower()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    errors, quality = [], []

    def note(metric, value):
        print(f"{metric}: {value}")

    note("run date", date.today().isoformat())
    note("scope", f"Centene, {args.state}, market {args.market}, network {args.network}")

    # C1
    index_urls, plans, index_grades = load_plans(errors)
    c1 = list({tuple(p.values()): p for p in plans}.values())
    files = Counter(p["rate_file_url"] for p in c1)
    write_csv(out / "c1_index_plans.csv", PLAN_COLUMNS, c1)
    for url, metrics in index_grades.items():
        quality += [{"file": url, "file_type": "index", **m} for m in metrics]
    note("index files found / parsed", f"{len(index_urls)} / {len(index_grades)}")
    note("plan entries / duplicates removed / plans", f"{len(plans):,} / {len(plans) - len(c1):,} / {len(c1):,}")
    note("unique rate files", len(files))
    note("index file quality scores", "; ".join(f"{Path(unquote(urlsplit(u).path)).name} {m[-1]['score']}"
                                                for u, m in index_grades.items()))

    # C2
    scoped = [p for p in c1 if p["market_type"] == args.market and state(p) == args.state]
    for p in scoped:
        p["network_type"], p["reason"] = network_type(p["plan_name"])
    target = [p for p in scoped if p["network_type"] == args.network]
    write_csv(out / f"c2_{net}_plans.csv", [*PLAN_COLUMNS, "network_type", "reason"], target)
    urls = list(dict.fromkeys(p["rate_file_url"] for p in target if p["rate_file_url"]))
    note(f"{args.state} {args.market} plans by network type",
         dict(Counter(p["network_type"] for p in scoped).most_common()))
    note(f"{args.network} plans / rate files", f"{len(target)} / {len(urls)}: {' '.join(urls)}")

    if not args.index_only and urls:
        # C4
        groups, tin_names, graded = {}, {}, []
        with ProcessPoolExecutor(min(len(urls), os.cpu_count() or 1)) as pool:
            jobs = [pool.submit(process_rate_file, u) for u in urls] if len(urls) > 1 else None
            for i, url in enumerate(urls):
                try:
                    file_groups, file_names, grader, stats = jobs[i].result() if jobs else process_rate_file(url)
                except FetchError as err:
                    errors.append(str(err))
                    continue
                merge(groups, file_groups)
                for tin, more in file_names.items():
                    tin_names.setdefault(tin, set()).update(more)
                graded.append((url, grader, stats["header"], file_groups))
                note(f"prices extracted ({Path(unquote(urlsplit(url).path)).name})",
                     f"{stats['prices']:,} prices -> {stats['rows']:,} rows x {stats['columns']} columns")
                if stats["unresolved_provider_references"]:
                    errors.append(f"{url}: rates point to undefined provider references "
                                  f"{stats['unresolved_provider_references'][:20]} (kept with blank NPI/TIN)")
                if stats["unknown_top_level_arrays"]:
                    errors.append(f"{url}: unknown top-level arrays not extracted: {stats['unknown_top_level_arrays']}")

        # C5
        npis = {npi for npi, _, _, _ in groups}
        print(f"Looking up {len(npis):,} NPIs in the NPI Registry...")
        providers = lookup_all(npis)
        errors += [f"NPI {n}: {p['note']}" for n, p in providers.items() if p.get("failed")]
        xwalk, spec_weights, spec_proc = load_supplement(args.supplement)
        try:
            nucc = load_nucc()
        except FetchError as err:
            errors.append(f"NUCC taxonomy: {err}")
            nucc = {}
        for p in providers.values():
            p["spec_code"], p["specialty"], p["specialty_source"] = specialty(p, xwalk, nucc)
        individuals = {n for n, p in providers.items() if p.get("npi_type", "").startswith("Type 1")}
        print("Streaming the CMS Doctors & Clinicians national file...")
        try:
            dac, cms_sites = cms_practice_sites(individuals)
        except (FetchError, ValueError, KeyError, StopIteration) as err:
            errors.append(f"CMS Doctors & Clinicians: {err!r}")
            dac, cms_sites = {}, set()
        states = {loc["state"] for p in providers.values() for loc in p.get("locations", [])} | {args.state}
        territories = {"AS", "GU", "MP", "PR", "VI"}
        irs = irs_organizations({tin for _, tin, _, _ in groups if tin}, states & (US_STATES - territories), errors)
        pairs = {(tin, npi) for npi, tin, _, _ in groups}
        links, tin_links = link(tin_names, pairs, providers, dac, irs, args.state, cms_sites)
        by_pair = {}
        for row in links:
            by_pair.setdefault((row["TIN"], row["NPI"]), []).append(row)
        removed = Counter()
        rated, rows = [], []
        medians = service_medians(groups)
        for row in rate_rows(groups, providers, by_pair, medians):
            reason = ("invalid NPI" if not npi_valid(row["NPI"]) else "no usable rate" if not row["Rate"]
                      else None)
            if reason:
                removed[reason] += 1
                continue
            rated.append(row)
            if row["Address"]:
                rows.append(row)
            else:
                removed["no address in any source"] += 1

        c4 = unique(rated, ["NPI", "TIN", "Billing Code"])
        write_csv(out / f"c4_{net}_rates.csv", ["NPI", "TIN", "Billing Code", "Rate"], c4)
        c5 = unique(rows, PROVIDER_RATE_COLUMNS)
        write_csv(out / f"c5_{net}_rates_by_address.csv", [*PROVIDER_RATE_COLUMNS, "Address Confidence"], c5)
        c6 = unique(rows, ["NPI", "TIN", "Address", "Billing Code"])
        write_xlsx(out / f"1_{net}_rate_file.xlsx", PROVIDER_RATE_COLUMNS, c6, numeric={"Rate"})

        # C3
        provider_spec = {n: p["spec_code"] for n, p in providers.items() if p.get("spec_code")}
        for url, grader, header, file_groups in graded:
            metrics = grader.finish(header, file_groups, provider_spec, spec_weights, spec_proc, providers)
            quality += [{"file": url, "file_type": "in-network rate file", **m} for m in metrics]
            note(f"rate file quality score ({Path(unquote(urlsplit(url).path)).name})", metrics[-1]["score"])

        # coverage is measured over real providers: valid NPIs, and TINs billing with at least one
        valid = {n for n in npis if npi_valid(n)}
        tins = {tin for tin, npi in pairs if npi in valid}
        verified = [r for r in links + tin_links if r["Address"] and r["Confidence"] in ("high", "medium")]
        note("C4 rows (NPI x TIN x billing code)", f"{len(c4):,}")
        note("C5 rows (NPI x TIN x address x billing code)", f"{len(c5):,}")
        note("C6 rows (NPI x TIN x address x billing code)", f"{len(c6):,}")
        note("C6 address confidence", dict(Counter(r["Address Confidence"] for r in c6).most_common()))
        note("rate file rows removed as invalid", dict(removed) or 0)
        note("how each rate was chosen", dict(Counter(
            "tie: lower median" if r["Rate Choice"].startswith("tie") else r["Rate Choice"].split(" (")[0]
            for r in c4).most_common()))
        note("NPI+TIN+code keys with several rates (one picked, see Rate Choice)",
             f"{sum(1 for g in groups.values() if len({r for r, _ in g['rates']}) > 1):,} of {len(groups):,}")
        outliers = sum(1 for g in groups.values() for r, svc in g["rates"] if is_outlier(r, medians.get(svc)))
        note("rates excluded ($0 or negative / expired / unreadable / outlier over 10x or under 1/10 of typical)",
             " / ".join(f"{sum(g[k] for g in groups.values()):,}" for k in ("zero", "expired", "unreadable"))
             + f" / {outliers:,}")
        note("unique NPIs / TINs", f"{len(valid):,} valid NPIs (+{len(npis) - len(valid)} invalid) / {len(tins):,} TINs")
        note("NPI types", dict(Counter(p.get("npi_type") or "not found" for p in providers.values())))
        note("specialty source", dict(Counter(p.get("specialty_source") or "none" for p in providers.values())))
        note("address sources", f"NPI Registry {len(providers):,} NPIs; CMS Doctors & Clinicians {len(dac):,} of "
                                f"{len(individuals):,} individuals and {len(cms_sites):,} active practice sites; "
                                f"IRS exempt organizations {len(irs):,} TINs")
        note("address links by confidence", dict(Counter(r["Confidence"] for r in links).most_common()))
        note("NPIs linked to a verified address (high or medium)", share({r["NPI"] for r in verified} & valid, valid))
        note("TINs linked to a verified address (high or medium)", share({r["TIN"] for r in verified} & tins, tins))
        note("NPIs / TINs with any address",
             f"{share({r['NPI'] for r in links if r['Address']} & valid, valid)} / "
             f"{share({r['TIN'] for r in links + tin_links if r['Address']} & tins, tins)}")

    (out / "c3_file_quality.md").write_text(quality_report(quality), encoding="utf-8")
    note("errors", len(errors))
    for e in errors:
        print(f"ERROR: {e}", file=sys.stderr)
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
